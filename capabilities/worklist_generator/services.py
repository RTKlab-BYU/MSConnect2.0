from __future__ import annotations

import re
import shutil
import tempfile
import uuid
from copy import copy
from pathlib import Path

from django.conf import settings
from openpyxl import load_workbook
from openpyxl.styles import Border, Font, Protection, Side

from .engine import run_worklist


ALLOWED_SUFFIXES = {".xlsx", ".xlsm"}
SAFE_FILENAME = re.compile(r"^[A-Za-z0-9._-]+$")
ASSETS_ROOT = Path(__file__).resolve().parent / "assets"


def results_root() -> Path:
    root = Path(settings.MEDIA_ROOT) / "worklist_generator_results"
    root.mkdir(parents=True, exist_ok=True)
    return root


def _safe_stem(value: str, fallback: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip()).strip("._-")
    return stem or fallback


def _job_directory() -> tuple[uuid.UUID, Path]:
    job_id = uuid.uuid4()
    job_dir = results_root() / str(job_id)
    job_dir.mkdir(parents=True, exist_ok=False)
    return job_id, job_dir


def _save_upload(uploaded_file, upload_dir: Path) -> Path:
    original_name = Path(uploaded_file.name or "workbook.xlsx").name
    suffix = Path(original_name).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise ValueError("Upload an Excel workbook with an .xlsx or .xlsm extension.")

    upload_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(mode="wb", suffix=suffix, dir=upload_dir, delete=False) as destination:
        for chunk in uploaded_file.chunks():
            destination.write(chunk)
        return Path(destination.name)


def _file_result(job_id: uuid.UUID, kind: str, path: Path) -> dict:
    return {"kind": kind, "name": path.name, "url_name": path.name}


def _metadata_groups(metadata_path: Path) -> tuple[str, list[str]]:
    try:
        workbook = load_workbook(metadata_path, read_only=True, data_only=True, keep_vba=metadata_path.suffix.lower() == ".xlsm")
        worksheet = workbook["DataEntry"]
    except KeyError as exc:
        raise ValueError("The metadata workbook must contain a DataEntry sheet.") from exc

    project_name = str(worksheet.cell(row=1, column=1).value or "").strip()
    groups: list[str] = []
    for column in range(3, worksheet.max_column + 1):
        value = worksheet.cell(row=8, column=column).value
        if value is None or not str(value).strip():
            break
        groups.append(str(value).strip())

    if not project_name:
        raise ValueError("The metadata workbook is missing a project name in DataEntry!A1.")
    if not groups:
        raise ValueError("The metadata workbook must define at least one condition in DataEntry row 8.")
    return project_name, groups


def generate_metadata_template(name: str, description: str, groups: list[str]) -> dict:
    """Create the CLI-compatible metadata workbook for WLG stage 1."""
    if not isinstance(groups, list):
        raise ValueError("Conditions must be provided as a list.")
    clean_name = name.strip()
    clean_description = description.strip()
    clean_groups = [group.strip() for group in groups if group.strip()]
    if not clean_name:
        raise ValueError("Enter a project name.")
    if not clean_groups:
        raise ValueError("Enter at least one condition name.")
    if len(clean_groups) > 50:
        raise ValueError("A maximum of 50 conditions can be included in one workbook.")
    if len(set(clean_groups)) != len(clean_groups):
        raise ValueError("Condition names must be unique.")

    job_id, job_dir = _job_directory()
    target = job_dir / f"{_safe_stem(clean_name, 'worklist')}_metadata.xlsm"
    source = ASSETS_ROOT / "metadataTemplate8.xlsm"
    try:
        shutil.copy2(source, target)
        workbook = load_workbook(target, keep_vba=True)
        worksheet = workbook["DataEntry"]
        worksheet.cell(row=1, column=1).value = clean_name
        worksheet.cell(row=1, column=2).value = clean_description
        format_pattern = []
        row = 9
        while True:
            cell = worksheet.cell(row=row, column=3)
            if cell.value == "END_OF_FILE":
                break
            if row > 200:
                raise ValueError("The metadata template is missing its END_OF_FILE marker.")
            format_pattern.append(cell)
            row += 1

        style_attrs = [
            {
                "font": copy(cell.font),
                "border": Border(
                    left=Side(style="thin", color="000000"),
                    right=Side(style="thin", color="000000"),
                    top=Side(style="thin", color="000000"),
                    bottom=Side(style="thin", color="000000"),
                ),
                "fill": copy(cell.fill),
                "number_format": cell.number_format,
                "protection": copy(cell.protection),
                "alignment": copy(cell.alignment),
            }
            for cell in format_pattern
        ]

        for index, group in enumerate(clean_groups, start=3):
            worksheet.cell(row=8, column=index).value = group
            worksheet.cell(row=8, column=index).font = Font(size=20, bold=True)
            for offset, style in enumerate(style_attrs):
                target_cell = worksheet.cell(row=9 + offset, column=index)
                target_cell.font = style["font"]
                target_cell.border = style["border"]
                target_cell.fill = style["fill"]
                target_cell.number_format = style["number_format"]
                target_cell.alignment = style["alignment"]
                target_cell.protection = Protection(locked=False)
        worksheet.protection.sheet = True
        workbook.save(target)
    except Exception:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise

    return {"job_id": str(job_id), "files": [_file_result(job_id, "metadata", target)]}


def generate_worklist_template(uploaded_file) -> dict:
    """Create the CLI-compatible plate-layout workbook for WLG stage 2."""
    job_id, job_dir = _job_directory()
    input_path: Path | None = None
    try:
        input_path = _save_upload(uploaded_file, Path(settings.MEDIA_ROOT) / "worklist_generator_uploads")
        project_name, groups = _metadata_groups(input_path)
        target = job_dir / f"{_safe_stem(project_name, 'worklist')}_worklist.xlsx"
        shutil.copy2(ASSETS_ROOT / "worklist_template_blank.xlsx", target)
        workbook = load_workbook(target)
        worksheet = workbook["User"]
        worksheet.cell(row=3, column=2).value = project_name
        for index, group in enumerate(groups, start=2):
            worksheet.cell(row=index, column=31).value = "Condition"
            worksheet.cell(row=index, column=32).value = group
        workbook.save(target)
    except Exception:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise
    finally:
        if input_path is not None:
            input_path.unlink(missing_ok=True)

    return {"job_id": str(job_id), "files": [_file_result(job_id, "worklist", target)]}


def generate_worklists(uploaded_file) -> dict:
    """Persist one generated pair in an isolated, downloadable job directory."""
    job_id, job_dir = _job_directory()
    input_path: Path | None = None
    try:
        input_path = _save_upload(uploaded_file, Path(settings.MEDIA_ROOT) / "worklist_generator_uploads")
        ms_path, lc_path = (Path(path) for path in run_worklist(input_path, output_dir=job_dir))
    except Exception:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise
    finally:
        if input_path is not None:
            input_path.unlink(missing_ok=True)

    files = [
        _file_result(job_id, "ms", ms_path),
        _file_result(job_id, "lc", lc_path),
    ]
    return {"job_id": str(job_id), "files": files}


def result_path(job_id: uuid.UUID, filename: str) -> Path | None:
    if not SAFE_FILENAME.fullmatch(filename) or Path(filename).suffix.lower() not in {".csv", ".xlsx", ".xlsm"}:
        return None
    root = (results_root() / str(job_id)).resolve()
    path = (root / filename).resolve()
    if path.parent != root or not path.is_file():
        return None
    return path
